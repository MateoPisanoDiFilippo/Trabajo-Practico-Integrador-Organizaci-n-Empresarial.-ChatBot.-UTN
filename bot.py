'''Trabajo Practico Integrador. Organización Empresarial.
ChatBot para alta de proveedores.
Heinzle Muriel Amancay. Comisión 8
Pisano Di Filippo Mateo Agustín. Comisión 8'''

import openpyxl
from datetime import date

ARCHIVO_EXCEL = "proveedores_murmat.xlsx"
HOJA = "proveedores"

# =========================================================
#  FUNCIONES DE DATOS Y VALIDACIÓN
# =========================================================

def cargar_proveedores():
    wb = openpyxl.load_workbook(ARCHIVO_EXCEL)
    hoja = wb[HOJA]
    proveedores = []
    for fila in hoja.iter_rows(min_row=2, values_only=True):
        if fila[0] is not None:
            proveedores.append({
                "id": fila[0],
                "nombre": fila[1],
                "cuit": str(fila[2]),
                "rubro": fila[3],
                "contacto": fila[4],
                "email": fila[5],
                "telefono": fila[6],
                "estado": fila[7],
                "fecha_alta": fila[8]
            })
    return proveedores

def normalizar_cuit(valor):
    """Quita guiones y espacios de un CUIT para comparar internamente."""
    return str(valor).strip().replace("-", "").replace(" ", "")

def proveedor_existe_cuit(cuit):
    """Verifica si ya existe un proveedor con ese CUIT (comparando sin guiones)."""
    cuit_norm = normalizar_cuit(cuit)
    for p in cargar_proveedores():
        if normalizar_cuit(p["cuit"]) == cuit_norm:
            return True
    return False

def proveedor_existe_nombre(nombre, excluir_cuit=None):
    """Verifica si ya existe un proveedor con ese nombre (ignorando mayúsculas).
    excluir_cuit permite saltear el proveedor actual al modificar."""
    for p in cargar_proveedores():
        if excluir_cuit and normalizar_cuit(p["cuit"]) == normalizar_cuit(excluir_cuit):
            continue
        if p["nombre"].lower() == nombre.lower():
            return True
    return False

def validar_cuit(cuit):
    """Valida que el CUIT tenga el formato obligatorio XX-XXXXXXXX-X con guiones.
    Devuelve el CUIT con guiones si es válido, o None si no lo es."""
    cuit = cuit.strip()
    partes = cuit.split("-")
    if len(partes) != 3:
        return None
    prefijo, medio, sufijo = partes
    # Formato: 2 dígitos - 8 dígitos - 1 dígito
    if not (prefijo.isdigit() and medio.isdigit() and sufijo.isdigit()):
        return None
    if not (len(prefijo) == 2 and len(medio) == 8 and len(sufijo) == 1):
        return None
    return cuit  # devuelve el CUIT con guiones

def validar_email(email):
    if "@" in email and "." in email:
        return email
    return None

def validar_telefono(telefono):
    telefono = telefono.replace(" ", "").replace("-", "")
    if telefono.isdigit() and len(telefono) >= 8:
        return telefono
    return None

def guardar_proveedor(datos):
    wb = openpyxl.load_workbook(ARCHIVO_EXCEL)
    hoja = wb[HOJA]
    nuevo_id = len(cargar_proveedores()) + 1
    fecha_hoy = date.today().strftime("%d/%m/%Y")
    hoja.append([
        nuevo_id,
        datos["nombre"],
        datos["cuit"],
        datos["rubro"],
        datos["contacto"],
        datos["email"],
        datos["telefono"],
        "activo",
        fecha_hoy
    ])
    wb.save(ARCHIVO_EXCEL)

def reordenar_ids():
    """Renumera completamente los IDs reescribiendo el archivo desde cero."""
    proveedores = cargar_proveedores()

    wb = openpyxl.load_workbook(ARCHIVO_EXCEL)
    hoja = wb[HOJA]

    # borrar todas las filas excepto encabezado
    hoja.delete_rows(2, hoja.max_row)

    # volver a escribir con IDs ordenados
    for i, p in enumerate(proveedores, start=1):
        hoja.append([
            i,
            p["nombre"],
            p["cuit"],
            p["rubro"],
            p["contacto"],
            p["email"],
            p["telefono"],
            p["estado"],
            p["fecha_alta"]
        ])

    wb.save(ARCHIVO_EXCEL)

def buscar_cuit_por_nombre(nombre_busqueda):
    """Devuelve el CUIT del primer proveedor cuyo nombre contenga la búsqueda."""
    for p in cargar_proveedores():
        if nombre_busqueda.lower() in p['nombre'].lower():
            return p['cuit']
    print('\nNo se encontró ningún proveedor con ese nombre.')
    return None

def resolver_cuit_o_nombre(texto):
    """Recibe lo que ingresó el usuario (CUIT con guiones o nombre) y devuelve el CUIT,
    o None si no se encontró o es inválido. Muestra los mensajes de error correspondientes."""
    if texto.replace('-', '').replace(' ', '').isdigit():
        # el usuario ingresó algo numérico → lo tratamos como CUIT
        cuit = validar_cuit(texto)
        if not cuit:
            print("CUIT inválido. El formato debe ser XX-XXXXXXXX-X (ej: 20-12345678-9).")
            return None
        return cuit
    else:
        # el usuario ingresó texto → lo tratamos como nombre
        return buscar_cuit_por_nombre(texto)

def mostrar_proveedores():
    proveedores = cargar_proveedores()
    if len(proveedores) == 0:
        print('No hay proveedores cargados en el sistema.')
        return
    print('---Lista de Proveedores---')
    for p in proveedores:
        print(f"""
ID:           {p['id']}
Nombre:       {p['nombre']}
CUIT:         {p['cuit']}
Rubro:        {p['rubro']}
Contacto:     {p['contacto']}
Email:        {p['email']}
Teléfono:     {p['telefono']}
Estado:       {p['estado']}
Fecha de alta:{p['fecha_alta']}
---------------------------""")


# =========================================================
#  OPERACIONES DEL MENÚ
# =========================================================

def alta_proveedor():
    datos = {}

    # nombre
    while True:
        nombre = input("Nombre o razón social del proveedor: ").strip().title()
        if nombre.lower() == "cancelar":
            print("Operación cancelada.")
            return
        if len(nombre) == 0:
            print("Debe ingresar algo. Intentá de nuevo.")
            continue
        if proveedor_existe_nombre(nombre):
            print("Ya existe un proveedor con ese nombre. Operación cancelada.")
            return
        datos["nombre"] = nombre
        break

    # cuit
    while True:
        cuit_input = input("CUIT con guiones (formato XX-XXXXXXXX-X): ").strip()
        if cuit_input.lower() == "cancelar":
            print("Operación cancelada.")
            return
        cuit = validar_cuit(cuit_input)
        if not cuit:
            print("CUIT inválido. El formato debe ser XX-XXXXXXXX-X (ej: 20-12345678-9).")
            continue
        if proveedor_existe_cuit(cuit):
            print("Este CUIT ya existe en el sistema. Operación cancelada.")
            return
        datos["cuit"] = cuit
        break

    # rubro
    while True:
        rubro = input("Rubro: ").strip().title()
        if rubro.lower() == "cancelar":
            print("Operación cancelada.")
            return
        if len(rubro) == 0:
            print("Debe ingresar algo. Intentá de nuevo.")
            continue
        datos["rubro"] = rubro
        break

    # contacto
    while True:
        contacto = input("Nombre de la persona de contacto: ").strip().title()
        if contacto.lower() == "cancelar":
            print("Operación cancelada.")
            return
        if len(contacto) == 0:
            print("Debe ingresar algo. Intentá de nuevo.")
            continue
        datos["contacto"] = contacto
        break

    # email
    while True:
        email = input("Email de contacto: ").strip()
        if email.lower() == "cancelar":
            print("Operación cancelada.")
            return
        if not validar_email(email):
            print("Email inválido. Debe contener @ y un dominio.")
            continue
        datos["email"] = email
        break

    # telefono
    while True:
        telefono = input("Teléfono de contacto: ").strip()
        if telefono.lower() == "cancelar":
            print("Operación cancelada.")
            return
        if not validar_telefono(telefono):
            print("Teléfono inválido. Solo números, mínimo 8 dígitos.")
            continue
        datos["telefono"] = telefono
        break

    # resumen y confirmación
    print("\n--- Resumen del nuevo proveedor ---")
    print(f"  Nombre:   {datos['nombre']}")
    print(f"  CUIT:     {datos['cuit']}")
    print(f"  Rubro:    {datos['rubro']}")
    print(f"  Contacto: {datos['contacto']}")
    print(f"  Email:    {datos['email']}")
    print(f"  Teléfono: {datos['telefono']}")
    print("-----------------------------------")
    while True:
        confirmar = input("¿Confirmás el alta? (s/n): ").strip().lower()
        if confirmar == "s":
            guardar_proveedor(datos)
            print("Proveedor dado de alta exitosamente en el sistema.")
            return
        elif confirmar == "n":
            print("Operación cancelada.")
            return
        else:
            print("Ingresá 's' para confirmar o 'n' para cancelar.")


def dar_baja_proveedor(cuit):
    wb = openpyxl.load_workbook(ARCHIVO_EXCEL)
    hoja = wb[HOJA]
    for fila in hoja.iter_rows(min_row=2):
        if normalizar_cuit(fila[2].value) == normalizar_cuit(cuit):
            print("\n---Proveedor encontrado.")
            print(f"  Nombre:   {fila[1].value}")
            print(f"  CUIT:     {fila[2].value}")
            print(f"  Rubro:    {fila[3].value}")
            print(f"  Contacto: {fila[4].value}")
            print("---------------------------")
            while True:
                confirmar = input("¿Confirmás la baja? (s/n): ").strip().lower()
                if confirmar == "s":
                    hoja.delete_rows(fila[0].row)
                    wb.save(ARCHIVO_EXCEL)
                    reordenar_ids()  # renumera los IDs correlativos
                    print("Proveedor eliminado del sistema exitosamente.")
                    return
                elif confirmar == "n":
                    print("Operación cancelada.")
                    return
                else:
                    print("Ingresá 's' para confirmar o 'n' para cancelar.")
    print("No se encontró ningún proveedor con ese CUIT.")


def cambiar_estados_proveedor(cuit):
    wb = openpyxl.load_workbook(ARCHIVO_EXCEL)
    hoja = wb[HOJA]
    for fila in hoja.iter_rows(min_row=2):
        if normalizar_cuit(fila[2].value) == normalizar_cuit(cuit):
            estado_actual = fila[7].value
            nuevo_estado = "inactivo" if estado_actual == "activo" else "activo"
            print("\n---Proveedor encontrado.")
            print(f"  Nombre:   {fila[1].value}")
            print(f"  CUIT:     {fila[2].value}")
            print(f"  Rubro:    {fila[3].value}")
            print(f"  Contacto: {fila[4].value}")
            print(f"  Estado:   {fila[7].value}")
            print("---------------------------")
            while True:
                confirmar = input("¿Confirmás la acción? (s/n): ").strip().lower()
                if confirmar == "s":
                    fila[7].value = nuevo_estado
                    wb.save(ARCHIVO_EXCEL)
                    print(f"\nEstado cambiado a '{nuevo_estado}' exitosamente.")
                    return
                elif confirmar == "n":
                    print("Operación cancelada.")
                    return
                else:
                    print("Ingresá 's' para confirmar o 'n' para cancelar.")
    print("No se encontró ningún proveedor con ese CUIT.")


def ver_proveedor():
    while True:
        cuit_nombre = input('\nIngresá el CUIT o nombre del proveedor a buscar: ').strip()
        if cuit_nombre.lower() == 'cancelar':
            print('\nOperación cancelada.')
            return
        cuit = resolver_cuit_o_nombre(cuit_nombre)
        if cuit is None:
            continue  # mensaje de error ya fue mostrado, pedimos de nuevo
        break

    for p in cargar_proveedores():
        if normalizar_cuit(p["cuit"]) == normalizar_cuit(cuit):
            print(f"""
ID:           {p['id']}
Nombre:       {p['nombre']}
CUIT:         {p['cuit']}
Rubro:        {p['rubro']}
Contacto:     {p['contacto']}
Email:        {p['email']}
Teléfono:     {p['telefono']}
Estado:       {p['estado']}
Fecha de alta:{p['fecha_alta']}
""")
            return
    print('\nNo se encontró el proveedor.')


def modificar_proveedor(cuit):
    wb = openpyxl.load_workbook(ARCHIVO_EXCEL)
    hoja = wb[HOJA]
    for fila in hoja.iter_rows(min_row=2):
        if normalizar_cuit(fila[2].value) == normalizar_cuit(cuit):
            print('\n---Proveedor encontrado.')
            print(f'  Nombre:   {fila[1].value}')
            print(f'  CUIT:     {fila[2].value}')
            print(f'  Rubro:    {fila[3].value}')
            print(f'  Contacto: {fila[4].value}')
            print(f'  Email:    {fila[5].value}')
            print(f'  Teléfono: {fila[6].value}')
            print('\nDejá vacío y presioná Enter para conservar el valor actual.')
            print('Escribí "cancelar" en cualquier momento para volver al menú.\n')

            # nombre
            while True:
                nombre = input(f'Nuevo nombre [{fila[1].value}]: ').strip()
                if nombre.lower() == "cancelar":
                    print('\nOperación cancelada.')
                    return
                if not nombre:
                    break
                if proveedor_existe_nombre(nombre.title(), excluir_cuit=cuit):
                    print("Ya existe otro proveedor con ese nombre. Intentá de nuevo.")
                    continue
                fila[1].value = nombre.title()
                break

            # cuit
            while True:
                cuit_nuevo = input(f'Nuevo CUIT [{fila[2].value}]: ').strip()
                if cuit_nuevo.lower() == "cancelar":
                    print('\nOperación cancelada.')
                    return
                if not cuit_nuevo:
                    break
                cuit_validado = validar_cuit(cuit_nuevo)
                if not cuit_validado:
                    print("CUIT inválido. El formato debe ser XX-XXXXXXXX-X (ej: 20-12345678-9).")
                    continue
                if normalizar_cuit(cuit_validado) != normalizar_cuit(cuit) and proveedor_existe_cuit(cuit_validado):
                    print("Ya existe otro proveedor con ese CUIT. Intentá de nuevo.")
                    continue
                fila[2].value = cuit_validado
                break

            # rubro
            rubro = input(f"Nuevo rubro [{fila[3].value}]: ").strip().title()
            if rubro.lower() == "cancelar":
                print('\nOperación cancelada.')
                return
            if rubro:
                fila[3].value = rubro

            # contacto
            contacto = input(f"Nuevo contacto [{fila[4].value}]: ").strip().title()
            if contacto.lower() == "cancelar":
                print('\nOperación cancelada.')
                return
            if contacto:
                fila[4].value = contacto

            # email
            while True:
                email = input(f"Nuevo email [{fila[5].value}]: ").strip()
                if email.lower() == "cancelar":
                    print('\nOperación cancelada.')
                    return
                if not email:
                    break
                if validar_email(email):
                    fila[5].value = email
                    break
                print("Email inválido. Debe contener @ y un dominio.")

            # telefono
            while True:
                telefono = input(f"Nuevo teléfono [{fila[6].value}]: ").strip()
                if telefono.lower() == "cancelar":
                    print('\nOperación cancelada.')
                    return
                if not telefono:
                    break
                telefono_validado = validar_telefono(telefono)
                if telefono_validado:
                    fila[6].value = telefono_validado
                    break
                print("Teléfono inválido. Solo números, mínimo 8 dígitos.")

            # confirmación
            while True:
                confirmar = input("\n¿Confirmás los cambios? (s/n): ").strip().lower()
                if confirmar == "s":
                    wb.save(ARCHIVO_EXCEL)
                    print("\nProveedor actualizado exitosamente.")
                    return
                elif confirmar == "n":
                    print("\nOperación cancelada.")
                    return
                else:
                    print('Ingresá "s" o "n".')
    print("\nNo se encontró ningún proveedor con ese CUIT.")


# =========================================================
#  MENÚ PRINCIPAL
# =========================================================

def menu():
    print()
    print("-" * 45)
    print("MURMAT S.A. — Sistema de Alta de Proveedores")
    print("-" * 45)
    print("Escribí 'cancelar' en cualquier momento para volver al menú.\n")

    while True:
        print("\n¿Qué deseas hacer?")
        print("""
1. Dar de alta un proveedor
2. Dar de baja un proveedor
3. Cambiar estado de un proveedor (Activo/Inactivo)
4. Ver lista completa de proveedores
5. Buscar proveedor por nombre/CUIT
6. Modificar datos del proveedor
7. Salir
""")
        opcion = input("Elegí una opción:\n").strip()

        if opcion == '7':
            print('Programa finalizado.')
            break

        elif opcion == "1":
            alta_proveedor()

        elif opcion == "2":
            while True:
                cuit_nombre = input('\nIngresá el CUIT o nombre del proveedor a dar de baja: ').strip()
                if cuit_nombre.lower() == 'cancelar':
                    print('\nOperación cancelada.')
                    break
                cuit = resolver_cuit_o_nombre(cuit_nombre)
                if cuit is None:
                    continue
                dar_baja_proveedor(cuit)
                break

        elif opcion == "3":
            while True:
                cuit_nombre = input('\nIngresá el CUIT o el nombre del proveedor a cambiar su estado: ').strip()
                if cuit_nombre.lower() == 'cancelar':
                    print('\nOperación cancelada.')
                    break
                cuit = resolver_cuit_o_nombre(cuit_nombre)
                if cuit is None:
                    continue
                cambiar_estados_proveedor(cuit)
                break

        elif opcion == '4':
            mostrar_proveedores()

        elif opcion == '5':
            ver_proveedor()

        elif opcion == '6':
            while True:
                cuit_nombre = input('\nIngresá el CUIT o nombre del proveedor a modificar: ').strip()
                if cuit_nombre.lower() == 'cancelar':
                    print('\nOperación cancelada.')
                    break
                cuit = resolver_cuit_o_nombre(cuit_nombre)
                if cuit is None:
                    continue
                modificar_proveedor(cuit)
                break

        else:
            print('\nOpción inválida.\n')


menu()